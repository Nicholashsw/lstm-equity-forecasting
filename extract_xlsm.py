"""Extract the Bloomberg BDH export sheets from the workbook into per-sheet CSVs.

Usage:
    python extract_xlsm.py <workbook.xlsm> [outdir]

Default outdir is `data/raw/bloomberg/` inside the repo, which is what
`src/parse_bbg_csv.py --indir data/raw/bloomberg` expects. The workbook itself is
licensed Bloomberg data and is not committed.
"""
import csv
import sys
import time
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent

if len(sys.argv) < 2:
    sys.exit(__doc__)
SRC = Path(sys.argv[1]).expanduser()
OUT = Path(sys.argv[2]).expanduser() if len(sys.argv) > 2 else ROOT / "data" / "raw" / "bloomberg"
OUT.mkdir(parents=True, exist_ok=True)

wb = load_workbook(SRC, read_only=True, data_only=True)
names = [n for n in wb.sheetnames if n.startswith("BBG_V")]
print(len(names), "BBG sheets")
for n in names:
    t0 = time.time()
    ws = wb[n]
    with open(OUT / f"{n}.csv", "w", newline="") as f:
        w = csv.writer(f)
        for row in ws.iter_rows(values_only=True):
            w.writerow(["" if v is None else v for v in row])
    print(f"{n} done {time.time()-t0:.0f}s", flush=True)
print("ALL DONE")
